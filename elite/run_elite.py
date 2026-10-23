"""Builds elite.db, runs every report in analysis_elite.sql, prints a preview of each and
writes a formatted Excel report (sales_report_elite.xlsx) with a KPI dashboard, charts and heatmaps."""
import os
import re
import sqlite3
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

os.chdir(os.path.dirname(os.path.abspath(__file__)))

subprocess.run([sys.executable, "generate_seed_elite.py"], check=True)

if os.path.exists("elite.db"):
    os.remove("elite.db")
con = sqlite3.connect("elite.db")
for f in ("schema_elite.sql", "seed_elite.sql"):
    con.executescript(open(f).read())

sql = open("analysis_elite.sql").read()
con.executescript(sql.split("-- @E1 ")[0])  # create the views

parts = re.split(r"-- @(E\d+) (.*)\n", sql)[1:]
results = []
for i in range(0, len(parts), 3):
    qid, title, body = parts[i], parts[i + 1].strip(), parts[i + 2]
    query = "\n".join(l for l in body.splitlines() if not l.strip().startswith("--")).strip().rstrip(";")
    cur = con.execute(query)
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    results.append((qid, title, cols, rows))

    print(f"\n=== {qid}: {title} ({len(rows)} rows) ===")
    shown = rows[:15]
    widths = [max([len(str(c))] + [len(str(r[j])) for r in shown]) for j, c in enumerate(cols)]
    print("  ".join(str(c).ljust(w) for c, w in zip(cols, widths)))
    for r in shown:
        print("  ".join(str(v).ljust(w) for v, w in zip(r, widths)))
    if not rows:
        print("(no rows)")
    elif len(rows) > 15:
        print(f"... {len(rows) - 15} more rows")
con.close()

# ---------------------------------------------------------------- Excel report
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
wb = Workbook()
wb.remove(wb.active)


def sheet_name(qid, title):
    return re.sub(r"[\[\]:*?/\\]", "", f"{qid} {title}")[:31]


def write_sheet(qid, title, cols, rows):
    ws = wb.create_sheet(sheet_name(qid, title))
    ws["A1"] = f"{qid}: {title}"
    ws["A1"].font = Font(bold=True, size=13)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=3, column=j, value=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, (int, float)):
                if cols[j - 1].endswith("pct"):
                    cell.number_format = '0.0"%"'
                elif isinstance(v, float) or abs(v) >= 1000:
                    cell.number_format = "#,##0"
            if row[0] == "TOTAL":
                cell.font = Font(bold=True)
    for j, c in enumerate(cols, 1):
        longest = max([len(str(c))] + [len(f"{r[j - 1]:,}") if isinstance(r[j - 1], (int, float))
                                       else len(str(r[j - 1])) for r in rows])
        ws.column_dimensions[get_column_letter(j)].width = min(longest + 4, 42)
    ws.freeze_panes = "A4"
    return ws, 3 + len(rows)


sheets = {}
for qid, title, cols, rows in results:
    ws, last = write_sheet(qid, title, cols, rows)
    sheets[qid] = (ws, last, cols)

# Charts and heatmaps
ws, last, _ = sheets["E2"]
ch = LineChart()
ch.title, ch.height, ch.width = "Revenue by month: 2025 vs 2026", 8, 16
ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=3, max_row=last), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
ws.add_chart(ch, "G3")

ws, last, _ = sheets["E5"]
ch = LineChart()
ch.title, ch.height, ch.width = "Monthly revenue vs 3-month moving average", 8, 18
ch.add_data(Reference(ws, min_col=2, min_row=3, max_row=last), titles_from_data=True)
ch.add_data(Reference(ws, min_col=4, min_row=3, max_row=last), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
ws.add_chart(ch, "F3")

ws, last, _ = sheets["E3"]
ch = BarChart()
ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
ch.title, ch.height, ch.width = "Revenue by region and category", 8, 16
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=3, max_row=last - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last - 1))
ws.add_chart(ch, "H3")
ws.conditional_formatting.add(
    f"B4:E{last - 1}",
    ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))

ws, last, _ = sheets["E8"]
ws.conditional_formatting.add(
    f"B4:G{last}",
    ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="70AD47"))

# KPI dashboard (first sheet)
dash = wb.create_sheet("Dashboard", 0)
dash["A1"] = "Sales Analytics - Elite Report"
dash["A1"].font = Font(bold=True, size=18, color="1F3864")
dash["A2"] = "Jan 2025 - Sep 2026 | revenue is net of discounts and excludes returned items"
dash["A2"].font = Font(italic=True, color="666666")
kcols, krow = results[0][2], results[0][3][0]
labels = {"net_revenue": "Net revenue", "gross_profit": "Gross profit", "margin_pct": "Margin %",
          "orders": "Orders", "active_customers": "Active customers",
          "avg_order_value": "Avg order value", "return_rate_pct": "Return rate %"}
for j, (c, v) in enumerate(zip(kcols, krow)):
    col = 1 + j * 2
    lab = dash.cell(row=4, column=col, value=labels.get(c, c))
    lab.font = Font(bold=True, color="FFFFFF")
    lab.fill = HDR_FILL
    lab.alignment = Alignment(horizontal="center")
    val = dash.cell(row=5, column=col, value=v)
    val.font = Font(bold=True, size=16)
    val.alignment = Alignment(horizontal="center")
    val.number_format = '0.0"%"' if c.endswith("pct") else "#,##0"
    dash.column_dimensions[get_column_letter(col)].width = 20
    dash.column_dimensions[get_column_letter(col + 1)].width = 3
dash["A8"] = "Reports"
dash["A8"].font = Font(bold=True, size=13)
for i, (qid, title, _, _) in enumerate(results[1:], 9):
    cell = dash.cell(row=i, column=1, value=f"{qid}: {title}")
    cell.hyperlink = f"#'{sheet_name(qid, title)}'!A1"
    cell.font = Font(color="0563C1", underline="single")

wb.save("sales_report_elite.xlsx")
print("\nsales_report_elite.xlsx written")
